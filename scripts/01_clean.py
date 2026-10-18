"""01_clean.py - clean one month's LinkedIn city exports.

Place this file in: <project>/scripts/01_clean.py

Default input:  <project>/data/raw data/September/
Default output: <project>/data/clean data/September/

Matches the uploaded Barcelona manual cleaning:
  * Keep the eight base job columns and every skills/0, skills/1, ... column.
  * Preserve row order, sheet names, and all posting rows.
  * Strip leading/trailing whitespace in text cells (not internal whitespace).
  * Do NOT merge city files, deduplicate jobs, or detect new skills yet.

Usage (from ANY working directory):
    python scripts/01_clean.py
    python scripts/01_clean.py --month October
    python scripts/01_clean.py --month September --overwrite

Requires: pip install openpyxl
"""
from __future__ import annotations

import argparse
import os
import re
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


# =====================================
# SETTINGS - can also use --month
# =====================================
MONTH = "September"

# Keep exactly the columns from the manually cleaned Barcelona example.
# Source order is deliberately standardised so every city has the same schema.
BASE_COLUMNS = (
    "companyName",
    "employmentType",
    "industry",
    "jobDescription",
    "jobTitle",
    "location",
    "postedDate",
    "seniorityLevel",
)
SKILL_COLUMN = re.compile(r"skills/(\d+)", flags=re.IGNORECASE)
HEADER_FILL = PatternFill("solid", fgColor="17324D")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def clean_value(value):
    """Preserve data types and text interiors; trim only surrounding whitespace."""
    return value.strip() if isinstance(value, str) else value


def output_filename(raw_name: str) -> str:
    """raw_data_Barcelona.xlsx -> clean_data_Barcelona.xlsx"""
    stem = Path(raw_name).stem
    if stem.lower().startswith("raw_data_"):
        stem = "clean_data_" + stem[len("raw_data_"):]
    else:
        stem = stem + "_cleaned"
    return stem + ".xlsx"


def clean_workbook(input_path: Path, output_path: Path) -> tuple[int, int, int]:
    """Clean all nonempty worksheets, returning (rows, retained_columns, removed_columns).

    Workbook is first saved to a temporary output and atomically replaced to avoid
    leaving a partly written cleaned file if processing fails.
    """
    wb_in = load_workbook(input_path, read_only=True, data_only=True)
    wb_out = Workbook()
    wb_out.remove(wb_out.active)
    total_rows = 0
    max_kept = 0
    max_removed = 0
    nonempty_sheets = 0

    try:
        for ws_in in wb_in.worksheets:
            rows = ws_in.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None or not any(value is not None for value in header):
                continue  # Ignore an empty worksheet, not a data-bearing one.

            columns = [str(x).strip() if x is not None else "" for x in header]
            if len([x for x in columns if x]) != len(set(x for x in columns if x)):
                raise ValueError(f"Sheet {ws_in.title!r}: duplicate column names")

            positions = {name: i for i, name in enumerate(columns) if name}
            missing = [name for name in BASE_COLUMNS if name not in positions]
            if missing:
                raise ValueError(
                    f"Sheet {ws_in.title!r}: required columns missing: {', '.join(missing)}"
                )

            skills = sorted(
                (name for name in positions if SKILL_COLUMN.fullmatch(name)),
                key=lambda name: int(SKILL_COLUMN.fullmatch(name).group(1)),
            )
            retained = [*BASE_COLUMNS, *skills]
            ws_out = wb_out.create_sheet(title=ws_in.title)
            ws_out.append(retained)
            ws_out.freeze_panes = "A2"
            ws_out.sheet_view.showGridLines = False
            ws_out.auto_filter.ref = f"A1:{ws_out.cell(1, len(retained)).column_letter}1"
            for cell in ws_out[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(wrap_text=True)
            ws_out.row_dimensions[1].height = 26
            for j, name in enumerate(retained, 1):
                letter = ws_out.cell(1, j).column_letter
                ws_out.column_dimensions[letter].width = (
                    65 if name == "jobDescription" else
                    48 if name == "jobTitle" else
                    36 if name == "companyName" else
                    28 if name in {"location", "industry"} else 23
                )

            for output_row, raw_values in enumerate(rows, start=2):
                for column_index, field in enumerate(retained, start=1):
                    source_index = positions[field]
                    value = raw_values[source_index] if source_index < len(raw_values) else None
                    value = clean_value(value)
                    cell = ws_out.cell(row=output_row, column=column_index, value=value)
                    # Never interpret a scraped description starting '=' as an Excel formula.
                    if isinstance(value, str):
                        cell.data_type = "s"
                    elif isinstance(value, datetime):
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
                    elif isinstance(value, date):
                        cell.number_format = "yyyy-mm-dd"
                total_rows += 1

            nonempty_sheets += 1
            max_kept = max(max_kept, len(retained))
            max_removed = max(max_removed, len(columns) - len(retained))
            ws_out.auto_filter.ref = ws_out.dimensions

        if not nonempty_sheets:
            raise ValueError("No nonempty worksheets found")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        fd, temp_name = tempfile.mkstemp(prefix=".cleaning_", suffix=".xlsx", dir=output_path.parent)
        os.close(fd)
        try:
            wb_out.save(temp_name)
            os.replace(temp_name, output_path)
        finally:
            Path(temp_name).unlink(missing_ok=True)
    finally:
        wb_in.close()
        wb_out.close()

    return total_rows, max_kept, max_removed


def run(month: str, overwrite: bool = False) -> int:
    project_root = Path(__file__).resolve().parent.parent
    source_dir = project_root / "data" / "raw data" / month
    destination_dir = project_root / "data" / "clean data" / month

    print(f"Project root : {project_root}")
    print(f"Input folder : {source_dir}")
    print(f"Output folder: {destination_dir}\n")

    if not source_dir.is_dir():
        print(f"ERROR: Input folder not found: {source_dir}", file=sys.stderr)
        return 1

    files = sorted(
        (p for p in source_dir.rglob("*") if p.is_file()
         and p.suffix.lower() == ".xlsx" and not p.name.startswith("~$")),
        key=lambda p: str(p).casefold(),
    )
    if not files:
        print("ERROR: No .xlsx files found in the input folder.", file=sys.stderr)
        return 1

    processed = skipped = failed = total_rows = 0
    for file in files:
        # Preserve any city subfolders under the month as well.
        relative_subfolder = file.relative_to(source_dir).parent
        output = destination_dir / relative_subfolder / output_filename(file.name)
        if output.exists() and not overwrite:
            print(f"SKIP (already exists): {output.name}  [use --overwrite to replace]")
            skipped += 1
            continue
        try:
            rows, kept, removed = clean_workbook(file, output)
            print(f"OK: {file.name} -> {output.name} | rows={rows}, columns={kept}, removed={removed}")
            processed += 1
            total_rows += rows
        except Exception as exc:
            failed += 1
            print(f"ERROR: {file.name}: {exc}", file=sys.stderr)

    print(f"\nFinished: {processed} file(s) cleaned, {skipped} skipped, "
          f"{failed} failed; {total_rows} rows written.")
    return 1 if failed else 0


def main():
    parser = argparse.ArgumentParser(description="Clean all LinkedIn Excel files for a scraping month")
    parser.add_argument("--month", default=MONTH,
                        help=f"Folder name below raw data/ and clean data/ (default: {MONTH})")
    parser.add_argument("--overwrite", action="store_true",
                        help="Replace existing cleaned output files; raw files are never modified")
    args = parser.parse_args()
    return run(args.month, overwrite=args.overwrite)


if __name__ == "__main__":
    sys.exit(main())
