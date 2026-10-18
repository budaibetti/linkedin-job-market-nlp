"""02_merge.py - append cleaned LinkedIn city exports into one monthly table.

Location:
    <project>/scripts/02_merge.py

Input (default):
    <project>/data/clean data/September/clean_data_*.xlsx
Output (default):
    <project>/data/merged file/September/merged_2026-09.csv
    <project>/data/merged file/September/merged_2026-09.xlsx

Every row from every city is retained. This script does NOT deduplicate jobs,
assign job IDs, detect regex skills, or calculate AI mentions.

Usage from your project root:
    python scripts/02_merge.py
    python scripts/02_merge.py --month October --year 2026
    python scripts/02_merge.py --month September --year 2026 --overwrite

Requires:
    pip install openpyxl
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import tempfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill


# ============================================================
# SETTINGS
# ============================================================
MONTH = "September"
YEAR = 2026

# Extend this mapping if you scrape another city. The country here is the
# SEARCH MARKET (source file's city), not necessarily the job's actual location.
CITY_MARKETS = {
    "barcelona": ("Barcelona", "Spain"),
    "madrid": ("Madrid", "Spain"),
    "valencia": ("Valencia", "Spain"),
    "milan": ("Milan", "Italy"),
    "milano": ("Milan", "Italy"),
    "torino": ("Torino", "Italy"),
    "turin": ("Torino", "Italy"),
    "budapest": ("Budapest", "Hungary"),
}

SOURCE_TO_OUTPUT = {
    "companyName": "company_name",
    "employmentType": "employment_type",
    "industry": "industry",
    "jobDescription": "job_description",
    "jobTitle": "job_title",
    "location": "location",
    "postedDate": "posted_date",
    "seniorityLevel": "seniority_level",
}
SKILL_HEADER = re.compile(r"skills/(\d+)", re.IGNORECASE)

OUTPUT_COLUMNS = [
    "snapshot_month", "market_country", "search_city",
    "job_title", "company_name", "job_description", "location",
    "posted_date", "employment_type", "industry", "seniority_level",
    "skills", "skills_json", "source_file", "source_sheet", "source_row",
]

MONTH_NUMBERS = {
    "january": 1, "february": 2, "march": 3, "april": 4,
    "may": 5, "june": 6, "july": 7, "august": 8,
    "september": 9, "october": 10, "november": 11, "december": 12,
}


def search_city_and_market(filename: str) -> tuple[str, str]:
    """Extract the scraped search city from clean_data_Barcelona2.xlsx etc."""
    stem = Path(filename).stem
    if not stem.casefold().startswith("clean_data_"):
        raise ValueError(f"Filename must start with 'clean_data_': {filename}")
    city = stem[len("clean_data_"):].strip()
    # Remove common duplicate-export suffixes, e.g. Barcelona2, Torino (2).
    city = re.sub(r"\s*\(\d+\)$", "", city).strip()
    city = re.sub(r"[ _-]*\d+$", "", city).strip()
    mapped = CITY_MARKETS.get(city.casefold())
    if mapped is None:
        raise ValueError(
            f"Unknown city {city!r} in {filename}. "
            "Add it to CITY_MARKETS in 02_merge.py before merging."
        )
    return mapped


def as_text(value) -> str:
    """Normalise workbook values without modifying internal description text."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def read_sheet_rows(path: Path, source_dir: Path, snapshot_month: str):
    """Yield normalised rows from all nonempty sheets of one city workbook."""
    city, market = search_city_and_market(path.name)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in wb.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header_values = next(rows, None)
            if header_values is None or not any(v is not None for v in header_values):
                continue
            headers = [as_text(v) for v in header_values]
            if len([h for h in headers if h]) != len(set(h for h in headers if h)):
                raise ValueError(f"{path.name} / {sheet.title}: duplicate headers")
            columns = {name: i for i, name in enumerate(headers) if name}
            missing = [name for name in SOURCE_TO_OUTPUT if name not in columns]
            if missing:
                raise ValueError(
                    f"{path.name} / {sheet.title}: missing required columns: {missing}. "
                    "Run 01_clean.py on this export first."
                )
            skill_columns = sorted(
                ((int(m.group(1)), idx) for name, idx in columns.items()
                 if (m := SKILL_HEADER.fullmatch(name))),
                key=lambda p: p[0],
            )
            if not skill_columns:
                print(f"WARNING: {path.name} / {sheet.title} has no skills/N columns")

            for excel_row, cells in enumerate(rows, start=2):
                def value_at(index):
                    return as_text(cells[index]) if index < len(cells) else ""

                base = {
                    output_name: value_at(columns[source_name])
                    for source_name, output_name in SOURCE_TO_OUTPUT.items()
                }
                seen = set()
                original_skills = []
                for _, col_idx in skill_columns:
                    skill = value_at(col_idx)
                    if skill and skill.casefold() not in seen:
                        original_skills.append(skill)
                        seen.add(skill.casefold())

                row = {
                    "snapshot_month": snapshot_month,
                    "market_country": market,
                    "search_city": city,
                    **base,
                    # Human-readable field for inspection and compatibility.
                    "skills": ", ".join(original_skills),
                    # Lossless list for subsequent taxonomy matching.
                    "skills_json": json.dumps(original_skills, ensure_ascii=False),
                    "source_file": path.relative_to(source_dir).as_posix(),
                    "source_sheet": sheet.title,
                    "source_row": excel_row,
                }
                # Preserve all data-bearing job rows; note blank rows explicitly.
                if not any(row[key] for key in SOURCE_TO_OUTPUT.values()) and not original_skills:
                    print(f"WARNING: completely blank source row at {path.name} / {sheet.title} / {excel_row}; skipped")
                    continue
                yield row
    finally:
        wb.close()


def write_merged(files: list[Path], source_dir: Path, csv_path: Path,
                 xlsx_path: Path | None, snapshot_month: str) -> tuple[int, dict]:
    """Stream data into a canonical CSV, plus an optional review workbook."""
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    stats = Counter()
    city_counts = Counter()
    fd, tmp_csv = tempfile.mkstemp(prefix=".merging_", suffix=".csv", dir=csv_path.parent)
    os.close(fd)
    tmp_xlsx = None
    workbook = None
    try:
        if xlsx_path is not None:
            fd, tmp_xlsx = tempfile.mkstemp(prefix=".merging_", suffix=".xlsx", dir=xlsx_path.parent)
            os.close(fd)
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet("Jobs")
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = False
            sheet.column_dimensions["A"].width = 17
            sheet.column_dimensions["B"].width = 18
            sheet.column_dimensions["C"].width = 20
            sheet.column_dimensions["D"].width = 42
            sheet.column_dimensions["E"].width = 35
            sheet.column_dimensions["F"].width = 75
            header_cells = []
            for field in OUTPUT_COLUMNS:
                c = WriteOnlyCell(sheet, value=field)
                c.fill = PatternFill("solid", fgColor="17324D")
                c.font = Font(bold=True, color="FFFFFF")
                c.alignment = Alignment(wrap_text=True)
                header_cells.append(c)
            sheet.append(header_cells)

        with open(tmp_csv, "w", encoding="utf-8-sig", newline="") as output:
            writer = csv.DictWriter(output, fieldnames=OUTPUT_COLUMNS)
            writer.writeheader()
            for path in files:
                current_count = 0
                for row in read_sheet_rows(path, source_dir, snapshot_month):
                    writer.writerow(row)
                    if workbook is not None:
                        # Excel can store only 1,048,576 worksheet rows including header.
                        if stats["rows"] >= 1_048_575:
                            raise ValueError("Excel row limit reached: rerun with --csv-only")
                        cells = []
                        for field in OUTPUT_COLUMNS:
                            c = WriteOnlyCell(sheet, value=row[field] if field != "source_row" else int(row[field]))
                            if isinstance(c.value, str):
                                c.data_type = "s"  # avoid interpreting scraped text as formulas
                            cells.append(c)
                        sheet.append(cells)
                    stats["rows"] += 1
                    current_count += 1
                    city_counts[row["search_city"]] += 1
                    if not row["skills_json"] or row["skills_json"] == "[]":
                        stats["no_original_skills"] += 1
                    if not all(row[k] for k in ["job_title", "company_name", "job_description", "location"]):
                        stats["incomplete_identity"] += 1
                print(f"OK: {path.relative_to(source_dir).as_posix()} -> {current_count} row(s)")

        if stats["rows"] == 0:
            raise ValueError("No job rows found in the input workbooks")
        if workbook is not None:
            workbook.save(tmp_xlsx)
        os.replace(tmp_csv, csv_path)
        if tmp_xlsx is not None:
            os.replace(tmp_xlsx, xlsx_path)
    finally:
        if workbook is not None:
            workbook.close()
        Path(tmp_csv).unlink(missing_ok=True)
        if tmp_xlsx is not None:
            Path(tmp_xlsx).unlink(missing_ok=True)

    return stats["rows"], {"stats": dict(stats), "cities": dict(city_counts)}


def run(month: str, year: int, overwrite: bool = False, csv_only: bool = False) -> int:
    normalized_month = month.strip().casefold()
    if normalized_month not in MONTH_NUMBERS:
        print(f"ERROR: invalid month {month!r}; use an English month name", file=sys.stderr)
        return 1
    if not 2000 <= year <= 2100:
        print("ERROR: year must be between 2000 and 2100", file=sys.stderr)
        return 1

    project_root = Path(__file__).resolve().parent.parent
    folder_month = normalized_month.capitalize()
    snapshot_month = f"{year:04d}-{MONTH_NUMBERS[normalized_month]:02d}"
    source_dir = project_root / "data" / "clean data" / folder_month
    destination_dir = project_root / "data" / "merged file" / folder_month
    csv_path = destination_dir / f"merged_{snapshot_month}.csv"
    xlsx_path = None if csv_only else destination_dir / f"merged_{snapshot_month}.xlsx"

    print(f"Project root  : {project_root}")
    print(f"Source folder : {source_dir}")
    print(f"Output folder : {destination_dir}")
    print(f"Scraping month: {snapshot_month} (NOT LinkedIn postedDate)\n")

    if not source_dir.is_dir():
        print("ERROR: cleaned month folder not found; run 01_clean.py first", file=sys.stderr)
        return 1
    files = sorted(
        (p for p in source_dir.rglob("*") if p.is_file()
         and p.suffix.lower() == ".xlsx" and not p.name.startswith("~$")
         and p.name.casefold().startswith("clean_data_")),
        key=lambda p: p.relative_to(source_dir).as_posix().casefold(),
    )
    if not files:
        print("ERROR: no clean_data_*.xlsx files found", file=sys.stderr)
        return 1
    existing_outputs = [p for p in [csv_path, xlsx_path] if p is not None and p.exists()]
    if existing_outputs and not overwrite:
        print("ERROR: output already exists. Use --overwrite after checking the old output:", file=sys.stderr)
        for p in existing_outputs:
            print(f"  {p}", file=sys.stderr)
        return 1

    try:
        count, report = write_merged(files, source_dir, csv_path, xlsx_path, snapshot_month)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    print(f"\nSUCCESS: {count} job rows appended from {len(files)} city workbook(s).")
    print(f"City breakdown: {report['cities']}")
    print(f"Rows without scraped skills: {report['stats'].get('no_original_skills', 0)}")
    print(f"Rows needing later identity review: {report['stats'].get('incomplete_identity', 0)}")
    print("No job-level deduplication or skill detection was performed.")
    print(f"CSV : {csv_path}")
    if xlsx_path is not None:
        print(f"Excel: {xlsx_path}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Append cleaned city exports into one monthly table")
    parser.add_argument("--month", default=MONTH, help=f"Month folder (default: {MONTH})")
    parser.add_argument("--year", type=int, default=YEAR, help=f"Scraping year (default: {YEAR})")
    parser.add_argument("--overwrite", action="store_true", help="Replace existing merged files")
    parser.add_argument("--csv-only", action="store_true", help="Write only CSV for larger datasets")
    args = parser.parse_args()
    return run(args.month, args.year, overwrite=args.overwrite, csv_only=args.csv_only)


if __name__ == "__main__":
    sys.exit(main())
