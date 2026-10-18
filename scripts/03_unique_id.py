"""03_unique_id.py - persistent numeric IDs and monthly deduplication.

Put this script in <project>/scripts/03_unique_id.py.

Input:   data/03_merged file/September/merged_2026-09.csv
Outputs: data/04_data_with_id/September/jobs_2026-09.csv (and .xlsx)
         data/04_data_with_id/September/duplicates_2026-09.csv
         data/04_data_with_id/September/identity_review_2026-09.csv
Registry: data/04_data_with_id/_registry/job_ids.sqlite

Job identity = exact job_title + company_name + job_description + location,
using text already stripped at the edges by Step 2. Matching does NOT ignore
case, punctuation, or internal whitespace.

Numeric job_id is persistent across months; numeric job_month_id identifies a
unique (job, scrape month) pair. The registry is essential: back it up and
NEVER delete/recreate it after processing multiple months.

Duplicates within a month are merged, including their scraped skills; jobs
observed in multiple months remain in each respective month's output.
Missing identity fields do NOT exclude a posting. Rows with sufficient remaining
information use a labelled partial identity, while poorly identified postings
receive a stable *within-source* fallback ID. All incomplete rows are included
in the main output AND copied into identity_review for checking. The complete
four-field fingerprint algorithm is unchanged to preserve existing numeric IDs.

Usage from project root:
    python scripts/03_unique_id.py
    python scripts/03_unique_id.py --month October --year 2026
    python scripts/03_unique_id.py --month September --year 2026 --overwrite
    python scripts/03_unique_id.py --csv-only

Dependencies: Python standard library; openpyxl for optional Excel preview.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import sqlite3
import sys
import tempfile
from collections import OrderedDict
from pathlib import Path


SCRIPT_VERSION = "3.2.0-incomplete-included"

MONTH = "September"
YEAR = 2026
MONTH_NUMBERS = {
    "january": 1, "february": 2, "march": 3, "april": 4,
    "may": 5, "june": 6, "july": 7, "august": 8,
    "september": 9, "october": 10, "november": 11, "december": 12,
}
IDENTITY_FIELDS = ("job_title", "company_name", "job_description", "location")
REQUIRED_FIELDS = (
    "snapshot_month", *IDENTITY_FIELDS, "skills", "skills_json",
    "source_file", "source_sheet", "source_row",
)
ADDED_FIELDS = (
    "job_id", "job_month_id", "same_month_source_count", "source_records_json",
    "identity_status", "identity_missing_fields",
)


def job_fingerprint(row: dict[str, str]) -> str:
    """Use the four fields, without ambiguous joining or lossy truncation."""
    payload = json.dumps(
        [row[name] for name in IDENTITY_FIELDS],
        ensure_ascii=False, separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def identity_for_row(row: dict[str, str], line_number: int, month: str) -> tuple[str, str, str]:
    """Return (fingerprint, identity_status, missing_fields).

    Complete matches keep the EXACT original hashing algorithm for registry
    compatibility. Partial matches are explicitly namespaced and are used only
    when the description and two other identifying fields are present. Sparse
    rows are keyed to their source record in that month, avoiding the dangerous
    assumption that two largely empty rows are the same posting.
    """
    missing = [field for field in IDENTITY_FIELDS if not row[field].strip()]
    if not missing:
        return job_fingerprint(row), "complete", ""

    missing_text = ", ".join(missing)
    present_count = len(IDENTITY_FIELDS) - len(missing)
    if row["job_description"].strip() and present_count >= 3:
        # If location is unknown, include the scrape market/search city so
        # identical text retrieved from different locations is not merged.
        market = [row.get("market_country", ""), row.get("search_city", "")] if "location" in missing else []
        payload = ["partial-v1", [row[field] for field in IDENTITY_FIELDS], market]
        encoded = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        return hashlib.sha256(encoded).hexdigest(), "partial", missing_text

    source = [row.get(field, "").strip() for field in ("source_file", "source_sheet", "source_row")]
    # If source coordinates are incomplete, tie this identity to the CSV line
    # as well. With complete coordinates, changing file order does not alter it.
    if not all(source):
        source.append(str(line_number))
    payload = ["source-fallback-v1", month, source, [row[field] for field in IDENTITY_FIELDS]]
    encoded = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest(), "source_fallback", missing_text


def parse_skills(row: dict[str, str], line_number: int) -> list[str]:
    """Read the lossless skill list; don't split names containing commas."""
    try:
        skills = json.loads(row["skills_json"])
    except (TypeError, ValueError) as exc:
        raise ValueError(f"CSV row {line_number}: malformed skills_json") from exc
    if not isinstance(skills, list) or any(not isinstance(s, str) for s in skills):
        raise ValueError(f"CSV row {line_number}: skills_json must be a list of strings")
    return [skill for s in skills if (skill := s.strip())]


def source_record(row: dict[str, str]) -> dict[str, str]:
    return {name: row[name] for name in ("source_file", "source_sheet", "source_row")}


def collect_month(source_path: Path, expected_month: str):
    """Collect all postings, including those with incomplete identities."""
    csv.field_size_limit(min(sys.maxsize, 2**31 - 1))
    records = OrderedDict()  # fingerprint -> row, merged skills, source references
    duplicates = []
    identity_review = []  # references to *included* jobs; attach IDs later
    with source_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        columns = reader.fieldnames or []
        if len(columns) != len(set(columns)):
            raise ValueError("Merged CSV contains duplicate column headers")
        missing = [name for name in REQUIRED_FIELDS if name not in columns]
        if missing:
            raise ValueError(f"Missing required fields: {missing}; rerun 02_merge.py")
        conflicting = [name for name in ADDED_FIELDS if name in columns]
        if conflicting:
            raise ValueError(f"Merged CSV already contains Step 3 fields: {conflicting}")

        for line_number, row in enumerate(reader, start=2):
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"CSV row {line_number}: incorrect number of columns")
            if row["snapshot_month"] != expected_month:
                raise ValueError(
                    f"CSV row {line_number}: snapshot_month {row['snapshot_month']!r} "
                    f"does not match {expected_month!r}"
                )
            fingerprint, status, missing_fields = identity_for_row(row, line_number, expected_month)
            skills = parse_skills(row, line_number)
            source = source_record(row)
            annotated = {
                **row, "identity_status": status,
                "identity_missing_fields": missing_fields,
            }
            if status != "complete":
                # Kept in both jobs and review. Duplicate source records each
                # get a review entry for traceability.
                identity_review.append({
                    **annotated, "_fingerprint": fingerprint,
                    "issue": f"{status}: missing {missing_fields}",
                })

            if fingerprint not in records:
                records[fingerprint] = {
                    "row": annotated, "skills": skills[:],
                    "seen": {skill.casefold() for skill in skills},
                    "sources": [source],
                }
            else:
                first = records[fingerprint]
                first["sources"].append(source)
                for skill in skills:
                    if skill.casefold() not in first["seen"]:
                        first["skills"].append(skill)
                        first["seen"].add(skill.casefold())
                duplicates.append({
                    **annotated, "_fingerprint": fingerprint,
                    "issue": "duplicate within snapshot month",
                    "retained_source_file": first["row"]["source_file"],
                    "retained_source_sheet": first["row"]["source_sheet"],
                    "retained_source_row": first["row"]["source_row"],
                })

    jobs = []
    for fingerprint, record in records.items():
        skills = record["skills"]
        jobs.append({
            "_fingerprint": fingerprint,
            **record["row"],
            "skills": ", ".join(skills),
            "skills_json": json.dumps(skills, ensure_ascii=False),
            "same_month_source_count": len(record["sources"]),
            "source_records_json": json.dumps(record["sources"], ensure_ascii=False),
        })
    return jobs, duplicates, identity_review, list(columns)


def create_tables(connection: sqlite3.Connection) -> None:
    connection.execute("PRAGMA foreign_keys=ON")
    connection.execute("PRAGMA busy_timeout=30000")
    connection.executescript("""
        CREATE TABLE IF NOT EXISTS jobs (
            job_id INTEGER PRIMARY KEY AUTOINCREMENT,
            fingerprint TEXT NOT NULL UNIQUE
        );
        CREATE TABLE IF NOT EXISTS job_months (
            job_month_id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id INTEGER NOT NULL,
            snapshot_month TEXT NOT NULL,
            UNIQUE(job_id, snapshot_month),
            FOREIGN KEY(job_id) REFERENCES jobs(job_id)
        );
    """)


def get_or_create_job_id(connection: sqlite3.Connection, fingerprint: str) -> int:
    result = connection.execute(
        "SELECT job_id FROM jobs WHERE fingerprint = ?", (fingerprint,)
    ).fetchone()
    if result is not None:
        return int(result[0])
    cursor = connection.execute(
        "INSERT INTO jobs (fingerprint) VALUES (?)", (fingerprint,)
    )
    return int(cursor.lastrowid)


def get_or_create_job_month_id(
    connection: sqlite3.Connection, job_id: int, month: str,
) -> int:
    result = connection.execute(
        "SELECT job_month_id FROM job_months WHERE job_id = ? AND snapshot_month = ?",
        (job_id, month),
    ).fetchone()
    if result is not None:
        return int(result[0])
    cursor = connection.execute(
        "INSERT INTO job_months (job_id, snapshot_month) VALUES (?, ?)",
        (job_id, month),
    )
    return int(cursor.lastrowid)


def assign_numeric_ids(connection, jobs, duplicates, review, month):
    """Attach numeric keys to EVERY job and to all audit records."""
    by_fingerprint = {}
    for row in jobs:
        fp = row.pop("_fingerprint")
        job_id = get_or_create_job_id(connection, fp)
        job_month_id = get_or_create_job_month_id(connection, job_id, month)
        row["job_id"] = job_id
        row["job_month_id"] = job_month_id
        by_fingerprint[fp] = (job_id, job_month_id)
    for rows in (duplicates, review):
        for row in rows:
            fp = row.pop("_fingerprint")
            row["job_id"], row["job_month_id"] = by_fingerprint[fp]
    jobs.sort(key=lambda row: row["job_id"])


def csv_to_temp(path: Path, fieldnames: list[str], rows: list[dict]) -> Path:
    fd, temp = tempfile.mkstemp(prefix=".step3_", suffix=".csv", dir=path.parent)
    os.close(fd)
    result = Path(temp)
    try:
        with result.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="raise")
            writer.writeheader()
            writer.writerows(rows)
        return result
    except Exception:
        result.unlink(missing_ok=True)
        raise


def excel_to_temp(path: Path, columns: list[str], rows: list[dict]) -> Path:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if len(rows) > 1_048_575:
        raise ValueError("Excel row limit exceeded; use --csv-only")
    fd, temp = tempfile.mkstemp(prefix=".step3_", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    result = Path(temp)
    workbook = Workbook(write_only=True)
    try:
        sheet = workbook.create_sheet("Jobs")
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for index, field in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = (
                17 if field in ("job_id", "job_month_id") else 28
            )
        headings = []
        for field in columns:
            cell = WriteOnlyCell(sheet, value=field)
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True)
            headings.append(cell)
        sheet.append(headings)
        for row in rows:
            cells = []
            for field in columns:
                value = row.get(field, "")
                if isinstance(value, str) and len(value) > 32767:
                    raise ValueError(
                        f"Column {field!r} exceeds Excel's 32767-character limit; "
                        "use --csv-only (CSV preserves full text)"
                    )
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # never interpret scraped text as an Excel formula
                cells.append(cell)
            sheet.append(cells)
        workbook.save(result)
        return result
    except Exception:
        result.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()


def run(month: str, year: int, overwrite: bool = False, csv_only: bool = False) -> int:
    month_key = month.strip().casefold()
    if month_key not in MONTH_NUMBERS:
        print(f"ERROR: unknown month {month!r}", file=sys.stderr)
        return 1
    if not 2000 <= year <= 2100:
        print("ERROR: invalid year", file=sys.stderr)
        return 1

    root = Path(__file__).resolve().parent.parent
    folder = month_key.capitalize()
    snapshot = f"{year:04d}-{MONTH_NUMBERS[month_key]:02d}"
    input_csv = root / "data" / "03_merged file" / folder / f"merged_{snapshot}.csv"
    processed_root = root / "data" / "04_data_with_id"
    output_dir = processed_root / folder
    registry = processed_root / "_registry" / "job_ids.sqlite"
    output_csv = output_dir / f"jobs_{snapshot}.csv"
    output_xlsx = output_dir / f"jobs_{snapshot}.xlsx"
    output_dup = output_dir / f"duplicates_{snapshot}.csv"
    output_review = output_dir / f"identity_review_{snapshot}.csv"
    targets = [output_csv, output_dup, output_review]
    if not csv_only:
        targets.append(output_xlsx)

    print(f"Script version: {SCRIPT_VERSION}")
    print(f"Script file   : {Path(__file__).resolve()}")
    print(f"Project root  : {root}")
    print(f"Input file    : {input_csv}")
    print(f"Output folder : {output_dir}")
    print(f"ID registry   : {registry}")
    if not input_csv.is_file():
        print("ERROR: merged CSV not found; run 02_merge.py first", file=sys.stderr)
        return 1
    existing = [path for path in targets if path.exists()]
    if existing and not overwrite:
        print("ERROR: output already exists; use --overwrite after checking:", file=sys.stderr)
        for path in existing:
            print(f"  {path}", file=sys.stderr)
        return 1

    # Prevent accidental ID reuse if someone deletes the registry but leaves
    # historical monthly outputs. Initial conversion of THIS month is allowed.
    if not registry.exists():
        # If Step 3 was previously run using the old folder, require users to
        # preserve its registry before switching to the new output location.
        legacy_root = root / "data" / "processed data"
        legacy_registry = legacy_root / "_registry" / "job_ids.sqlite"
        legacy_outputs = list(legacy_root.glob("*/jobs_*.csv"))
        if legacy_registry.exists() or legacy_outputs:
            print(
                "ERROR: previous Step 3 data exists under 'data/processed data', "
                "but no registry exists in 'data/04_data_with_id/_registry'. "
                "Copy the existing job_ids.sqlite into the new _registry folder "
                "before continuing. Do not start a new registry, as numeric IDs "
                "could change.",
                file=sys.stderr,
            )
            return 1
        other_months = [
            p for p in processed_root.glob("*/jobs_*.csv")
            if p.parent != output_dir
        ]
        if other_months:
            print(
                "ERROR: persistent ID registry is missing, but other months have "
                "processed outputs. Restore the registry from a backup, or "
                "rebuild ALL months in order into a fresh processed-data folder.",
                file=sys.stderr,
            )
            return 1

    staged = []
    connection = None
    try:
        jobs, duplicates, review, input_columns = collect_month(input_csv, snapshot)
        if not jobs:
            raise ValueError("Merged CSV contains no data rows")
        output_dir.mkdir(parents=True, exist_ok=True)
        registry.parent.mkdir(parents=True, exist_ok=True)
        connection = sqlite3.connect(registry, timeout=30)
        create_tables(connection)
        connection.execute("BEGIN IMMEDIATE")
        assign_numeric_ids(connection, jobs, duplicates, review, snapshot)

        columns = ["job_id", "job_month_id", *input_columns,
                   "identity_status", "identity_missing_fields",
                   "same_month_source_count", "source_records_json"]
        staged.append((csv_to_temp(output_csv, columns, jobs), output_csv))
        staged.append((csv_to_temp(output_dup, input_columns + [
            "job_id", "job_month_id", "issue", "retained_source_file",
            "retained_source_sheet", "retained_source_row",
            "identity_status", "identity_missing_fields",
        ], duplicates), output_dup))
        staged.append((csv_to_temp(output_review, input_columns + [
            "job_id", "job_month_id", "identity_status", "identity_missing_fields", "issue"
        ], review), output_review))
        if not csv_only:
            staged.append((excel_to_temp(output_xlsx, columns, jobs), output_xlsx))
        connection.commit()  # Persist registry BEFORE releasing updated monthly files.
        for temp_path, target in staged:
            os.replace(temp_path, target)
    except Exception as exc:
        if connection is not None:
            connection.rollback()
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    finally:
        if connection is not None:
            connection.close()
        for temp_path, _ in staged:
            temp_path.unlink(missing_ok=True)

    print(f"\nSUCCESS: {len(jobs)} unique job(s) included for {snapshot}")
    print(f"  Complete identities: {sum(row['identity_status'] == 'complete' for row in jobs)}")
    print(f"  Partial identities: {sum(row['identity_status'] == 'partial' for row in jobs)}")
    print(f"  Source fallbacks: {sum(row['identity_status'] == 'source_fallback' for row in jobs)}")
    print(f"Duplicate row(s) within month: {len(duplicates)}")
    print(f"Incomplete identity row(s) INCLUDED and flagged for review: {len(review)}")
    print("Numeric job_id and job_month_id are reused on reruns and future months.")
    print(f"Jobs CSV     : {output_csv}")
    if not csv_only:
        print(f"Jobs Excel   : {output_xlsx}")
    elif output_xlsx.exists():
        print(f"NOTE: existing Excel preview was not updated: {output_xlsx}")
    print(f"Duplicate log: {output_dup}")
    print(f"Review log   : {output_review}")
    print(f"KEEP/BACK UP : {registry} (needed to preserve numeric IDs)")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Assign persistent numeric job IDs and deduplicate within a scrape month"
    )
    parser.add_argument("--version", action="version", version=f"%(prog)s {SCRIPT_VERSION}")
    parser.add_argument("--month", default=MONTH, help=f"Month folder (default: {MONTH})")
    parser.add_argument("--year", type=int, default=YEAR, help=f"Scrape year (default: {YEAR})")
    parser.add_argument("--overwrite", action="store_true", help="Regenerate this month's output")
    parser.add_argument("--csv-only", action="store_true", help="Skip Excel preview")
    args = parser.parse_args()
    return run(args.month, args.year, overwrite=args.overwrite, csv_only=args.csv_only)


if __name__ == "__main__":
    sys.exit(main())
